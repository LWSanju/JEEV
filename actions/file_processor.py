from pathlib import Path

try:
    from memory.memory_manager import save_document_memory
except Exception:
    save_document_memory = None

import json
import csv
import mimetypes
import os
import time


# ============================================================
# JEEV — FILE INTELLIGENCE PROCESSOR
# FULL REPLACEMENT
# ============================================================


MAX_TEXT_CHARS = 30000
MAX_SHEET_ROWS = 500
MAX_CELL_CHARS = 500

IMAGE_MAX_BYTES = 20 * 1024 * 1024

# Maximum number of pages used by the visual fallback.
MAX_PDF_VISUAL_PAGES = 20

# PDF rendering quality.
PDF_RENDER_SCALE = 1.8

# Minimum amount of extracted text considered useful.
PDF_MIN_USEFUL_TEXT = 80

# Gemini direct PDF input limit.
MAX_DIRECT_PDF_BYTES = 50 * 1024 * 1024


# ============================================================
# EXTENSIONS
# ============================================================

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".py",
    ".pyw",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".cs",
    ".go",
    ".rs",
    ".php",
    ".rb",
    ".swift",
    ".kt",
    ".kts",
    ".html",
    ".htm",
    ".css",
    ".scss",
    ".sass",
    ".xml",
    ".yaml",
    ".yml",
    ".ini",
    ".cfg",
    ".conf",
    ".log",
    ".sql",
    ".sh",
    ".bat",
    ".ps1",
    ".env",
}

IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".bmp",
    ".gif",
    ".tif",
    ".tiff",
    ".ico",
}


# ============================================================
# GENERAL HELPERS
# ============================================================

def _safe_string(value):
    if value is None:
        return ""

    try:
        return str(value).strip()
    except Exception:
        return ""


def _get_file_model():
    """
    Gemini model used for JEEV file/PDF/vision analysis.

    You can override it with:

        set GEMINI_FILE_MODEL=gemini-3.6-flash

    """

    return os.getenv(
        "GEMINI_FILE_MODEL",
        "gemini-3.6-flash",
    )


def _find_gemini_client(player=None, client=None):
    """
    Locate the Gemini client.

    Priority:

        1. Explicit client
        2. player.client
        3. player.gemini_client
        4. player.ai_client
        5. player.genai_client
    """

    if client is not None:
        return client

    if player is None:
        return None

    possible_names = (
        "client",
        "gemini_client",
        "ai_client",
        "genai_client",
    )

    for name in possible_names:
        try:
            value = getattr(
                player,
                name,
                None,
            )

            if value is not None:
                return value

        except Exception:
            pass

    return None


def _limit_text(text):
    if not text:
        return "No readable content was found."

    text = str(text)

    if len(text) <= MAX_TEXT_CHARS:
        return text

    return (
        text[:MAX_TEXT_CHARS]
        + "\n\n"
        "[Content truncated because the file is very large.]"
    )


# ============================================================
# IMAGE INFORMATION
# ============================================================

def _image_info(path):
    try:
        from PIL import Image

        with Image.open(path) as image:
            return (
                f"Image format: {image.format}\n"
                f"Dimensions: {image.width} x {image.height}\n"
                f"Color mode: {image.mode}"
            )

    except ImportError:
        return (
            "Image detected. "
            "Install Pillow for image dimensions:\n"
            "python -m pip install Pillow"
        )

    except Exception as e:
        return (
            "Image detected, but metadata could not be read: "
            f"{e}"
        )


# ============================================================
# TEXT FILE
# ============================================================

def _read_text_file(path):
    try:
        return path.read_text(
            encoding="utf-8",
            errors="replace",
        )

    except Exception as e:
        return f"Could not read text file: {e}"


# ============================================================
# PDF — PYPDF TEXT EXTRACTION
# ============================================================

def _read_pdf_with_pypdf(path):
    try:
        import pypdf

    except ImportError as e:
        return (
            "",
            "pypdf is not installed.\n"
            "Install it with:\n"
            "python -m pip install pypdf\n\n"
            f"Import error: {e}",
        )

    try:
        reader = pypdf.PdfReader(str(path))

        pages = []

        for index, page in enumerate(reader.pages):

            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""

            if text.strip():
                pages.append(
                    f"\n--- Page {index + 1} ---\n"
                    f"{text.strip()}"
                )

        return "\n".join(pages), ""

    except Exception as e:
        return (
            "",
            f"pypdf processing failed: {e}",
        )


# ============================================================
# PDF — PYMUPDF TEXT EXTRACTION
# ============================================================

def _read_pdf_with_pymupdf(path):
    try:
        import fitz

    except ImportError as e:
        return (
            "",
            "PyMuPDF is not installed.\n"
            "Install it with:\n"
            "python -m pip install pymupdf\n\n"
            f"Import error: {e}",
        )

    document = None

    try:
        document = fitz.open(str(path))

        pages = []

        for index, page in enumerate(document):

            try:
                text = page.get_text("text") or ""
            except Exception:
                text = ""

            if text.strip():
                pages.append(
                    f"\n--- Page {index + 1} ---\n"
                    f"{text.strip()}"
                )

        return "\n".join(pages), ""

    except Exception as e:
        return (
            "",
            f"PyMuPDF processing failed: {e}",
        )

    finally:
        if document is not None:
            try:
                document.close()
            except Exception:
                pass


# ============================================================
# PDF — MAIN TEXT READER
# ============================================================

def _read_pdf(path):
    """
    Local text extraction.

    This is only a fallback.
    Gemini visual/native PDF analysis is preferred.
    """

    errors = []

    pymupdf_text, pymupdf_error = (
        _read_pdf_with_pymupdf(path)
    )

    if pymupdf_text.strip():
        return pymupdf_text

    if pymupdf_error:
        errors.append(pymupdf_error)

    pypdf_text, pypdf_error = (
        _read_pdf_with_pypdf(path)
    )

    if pypdf_text.strip():
        return pypdf_text

    if pypdf_error:
        errors.append(pypdf_error)

    if errors:
        return (
            "PDF text extraction could not produce "
            "readable text.\n\n"
            + "\n".join(errors)
            + "\n\n"
            "No selectable text was found."
        )

    return (
        "The PDF was opened, but no selectable text "
        "could be extracted.\n\n"
        "This may be a scanned or image-only PDF."
    )


# ============================================================
# PDF — DETECT SELECTABLE TEXT
# ============================================================

def _pdf_has_selectable_text(path):

    pymupdf_text, _ = (
        _read_pdf_with_pymupdf(path)
    )

    if len(pymupdf_text.strip()) >= PDF_MIN_USEFUL_TEXT:
        return True

    pypdf_text, _ = (
        _read_pdf_with_pypdf(path)
    )

    if len(pypdf_text.strip()) >= PDF_MIN_USEFUL_TEXT:
        return True

    return False


# ============================================================
# PDF — PAGE COUNT
# ============================================================

def _pdf_page_count(path):

    try:
        import fitz

        document = fitz.open(str(path))

        try:
            return len(document)
        finally:
            document.close()

    except Exception:

        try:
            import pypdf

            reader = pypdf.PdfReader(
                str(path)
            )

            return len(reader.pages)

        except Exception:
            return 0


# ============================================================
# PDF — RENDER PAGE
# ============================================================

def _render_pdf_page(
    path,
    page_number,
):

    try:
        import fitz

    except ImportError:
        raise RuntimeError(
            "Scanned PDF support requires PyMuPDF.\n"
            "Install it with:\n"
            "python -m pip install pymupdf"
        )

    document = None

    try:
        document = fitz.open(str(path))

        if (
            page_number < 0
            or page_number >= len(document)
        ):
            raise IndexError(
                f"PDF page {page_number + 1} "
                "does not exist."
            )

        page = document[page_number]

        matrix = fitz.Matrix(
            PDF_RENDER_SCALE,
            PDF_RENDER_SCALE,
        )

        pixmap = page.get_pixmap(
            matrix=matrix,
            alpha=False,
        )

        image_bytes = pixmap.tobytes(
            "png"
        )

        return image_bytes, "image/png"

    finally:

        if document is not None:

            try:
                document.close()
            except Exception:
                pass


# ============================================================
# PDF — GEMINI DIRECT PDF ANALYSIS
# ============================================================

def _analyze_pdf_direct(
    path,
    instruction,
    client,
):
    """
    PRIMARY PDF ANALYSIS.

    Sends the original PDF directly to Gemini.

    This allows Gemini to understand:

        - selectable text
        - scanned pages
        - CV layouts
        - images
        - tables
        - diagrams
        - multiple pages
        - visual formatting
    """

    if client is None:
        return None

    try:
        file_size = path.stat().st_size

    except Exception as e:
        print(
            "[FILE PROCESSOR] "
            f"Could not read PDF size: {e}"
        )

        return None

    if file_size <= 0:
        return (
            "PDF ANALYSIS FAILED\n"
            "===================\n"
            "The PDF file is empty."
        )

    if file_size > MAX_DIRECT_PDF_BYTES:
        print(
            "[FILE PROCESSOR] "
            "PDF exceeds direct Gemini limit. "
            "Using visual fallback."
        )

        return None

    user_instruction = (
        instruction or ""
    ).strip()

    if not user_instruction:

        user_instruction = (
            "Analyze this PDF carefully.\n\n"

            "If this is a CV or resume, provide a "
            "complete professional analysis including:\n"

            "1. Candidate name and profile\n"
            "2. Contact information visible in the document\n"
            "3. Career objective or summary\n"
            "4. Education\n"
            "5. Work experience\n"
            "6. Internships\n"
            "7. Skills\n"
            "8. Projects\n"
            "9. Certifications\n"
            "10. Achievements\n"
            "11. Languages\n"
            "12. Strengths\n"
            "13. Weaknesses or missing information\n"
            "14. Resume structure and formatting\n"
            "15. ATS/readability assessment\n"
            "16. Specific improvements\n"
            "17. Overall assessment\n\n"

            "Read the actual document, including text "
            "inside scanned pages or images.\n"

            "Do not say that PDF analysis is unsupported.\n"
            "Do not invent information.\n"

            "If something cannot be read, clearly mark "
            "it as unreadable."
        )

    prompt = (
        f"Analyze the PDF file '{path.name}'.\n\n"

        f"{user_instruction}\n\n"

        "IMPORTANT:\n"

        "Use the PDF itself as the source of truth. "
        "Inspect all pages. "

        "Do not assume that failure of local text "
        "extraction means the document cannot be analyzed. "

        "If the document is a scanned PDF, visually inspect "
        "the pages."
    )

    uploaded_file = None

    try:

        print(
            "[FILE PROCESSOR] "
            f"Uploading PDF to Gemini: {path.name}"
        )

        # ------------------------------------------------
        # Upload original PDF
        # ------------------------------------------------

        try:

            uploaded_file = client.files.upload(
                file=str(path),
            )

        except Exception as first_upload_error:

            print(
                "[FILE PROCESSOR] "
                "Normal PDF upload failed: "
                f"{first_upload_error}"
            )

            try:

                uploaded_file = client.files.upload(
                    file=str(path),
                    config={
                        "mime_type": "application/pdf"
                    },
                )

            except Exception as second_upload_error:

                print(
                    "[FILE PROCESSOR] "
                    "MIME PDF upload also failed: "
                    f"{second_upload_error}"
                )

                return None

        if uploaded_file is None:
            return None

        # ------------------------------------------------
        # Wait for Gemini file processing
        # ------------------------------------------------

        file_name = getattr(
            uploaded_file,
            "name",
            None,
        )

        if file_name:

            for _ in range(30):

                try:

                    current_file = (
                        client.files.get(
                            name=file_name
                        )
                    )

                    state = getattr(
                        current_file,
                        "state",
                        None,
                    )

                    state_name = str(
                        getattr(
                            state,
                            "name",
                            state,
                        )
                    ).upper()

                    print(
                        "[FILE PROCESSOR] "
                        f"Gemini PDF state: "
                        f"{state_name}"
                    )

                    if (
                        "PROCESSING"
                        not in state_name
                    ):
                        uploaded_file = (
                            current_file
                        )
                        break

                    if (
                        "FAILED"
                        in state_name
                    ):
                        print(
                            "[FILE PROCESSOR] "
                            "Gemini PDF processing failed."
                        )
                        return None

                    time.sleep(1)

                except Exception as state_error:

                    print(
                        "[FILE PROCESSOR] "
                        "Could not check PDF state: "
                        f"{state_error}"
                    )

                    break

        # ------------------------------------------------
        # Native Gemini PDF understanding
        # ------------------------------------------------

        print(
            "[FILE PROCESSOR] "
            "Sending PDF to Gemini for analysis..."
        )

        response = (
            client.models.generate_content(
                model=_get_file_model(),
                contents=[
                    prompt,
                    uploaded_file,
                ],
            )
        )

        result = getattr(
            response,
            "text",
            None,
        )

        if result and result.strip():

            print(
                "[FILE PROCESSOR] "
                "Gemini native PDF analysis succeeded."
            )

            return (
                "PDF ANALYSIS\n"
                "============\n"
                f"File: {path.name}\n\n"
                f"{result.strip()}"
            )

        print(
            "[FILE PROCESSOR] "
            "Gemini received PDF but returned "
            "no analysis."
        )

        return None

    except Exception as e:

        print(
            "[FILE PROCESSOR] "
            f"Native PDF analysis failed: {e}"
        )

        return None


# ============================================================
# PDF — GEMINI VISUAL PAGE ANALYSIS
# ============================================================

def _analyze_pdf_page(
    image_bytes,
    page_number,
    total_pages,
    instruction,
    client,
    file_name,
):

    if client is None:
        return (
            "Gemini client is not available "
            "for visual PDF analysis."
        )

    try:
        from google.genai import types

        image_part = (
            types.Part.from_bytes(
                data=image_bytes,
                mime_type="image/png",
            )
        )

        user_instruction = (
            instruction or ""
        ).strip()

        if not user_instruction:

            user_instruction = (
                "Analyze this document page carefully. "

                "Read all visible text. "

                "If this is a CV/resume, identify "
                "candidate details, education, experience, "
                "skills, projects, achievements, "
                "certifications and other useful information. "

                "Read text from the image itself even if "
                "the original PDF has no selectable text. "

                "Do not invent information."
            )

        prompt = (
            f"This is page {page_number} "
            f"of {total_pages} from "
            f"'{file_name}'.\n\n"

            f"{user_instruction}\n\n"

            "IMPORTANT:\n"

            "This page was rendered as an image because "
            "the PDF may be scanned or image-only. "

            "Visually inspect the entire page. "

            "Read names, dates, phone numbers, email "
            "addresses, headings, bullet points, tables "
            "and other visible information carefully. "

            "Return only information supported by "
            "the visible page."
        )

        response = (
            client.models.generate_content(
                model=_get_file_model(),
                contents=[
                    image_part,
                    prompt,
                ],
            )
        )

        result = getattr(
            response,
            "text",
            None,
        )

        if result and result.strip():
            return result.strip()

        return (
            f"Gemini returned no analysis "
            f"for page {page_number}."
        )

    except Exception as e:

        return (
            f"Visual analysis failed for "
            f"page {page_number}: {e}"
        )


# ============================================================
# PDF — SCANNED VISUAL FALLBACK
# ============================================================

def _analyze_scanned_pdf(
    path,
    instruction,
    client,
):

    if client is None:
        return (
            "PDF ANALYSIS FAILED\n"
            "===================\n"
            "Gemini client is not available."
        )

    total_pages = _pdf_page_count(
        path
    )

    if total_pages <= 0:
        return (
            "PDF ANALYSIS FAILED\n"
            "===================\n"
            "The PDF page count could not "
            "be determined."
        )

    pages_to_process = min(
        total_pages,
        MAX_PDF_VISUAL_PAGES,
    )

    print(
        "[FILE PROCESSOR] "
        f"Visual PDF fallback: "
        f"{pages_to_process}/{total_pages} pages."
    )

    analyses = []

    for page_index in range(
        pages_to_process
    ):

        page_number = page_index + 1

        try:

            print(
                "[FILE PROCESSOR] "
                f"Rendering PDF page "
                f"{page_number}/{pages_to_process}..."
            )

            image_bytes, _ = (
                _render_pdf_page(
                    path,
                    page_index,
                )
            )

            if len(image_bytes) > IMAGE_MAX_BYTES:

                analyses.append(
                    f"\n--- Page {page_number} ---\n"
                    "Page image was too large for "
                    "visual analysis."
                )

                continue

            analysis = (
                _analyze_pdf_page(
                    image_bytes=image_bytes,
                    page_number=page_number,
                    total_pages=total_pages,
                    instruction=instruction,
                    client=client,
                    file_name=path.name,
                )
            )

            analyses.append(
                f"\n--- Page {page_number} ---\n"
                f"{analysis}"
            )

            print(
                "[FILE PROCESSOR] "
                f"Page {page_number} analyzed."
            )

        except Exception as e:

            print(
                "[FILE PROCESSOR] "
                f"Page {page_number} failed: {e}"
            )

            analyses.append(
                f"\n--- Page {page_number} ---\n"
                f"Could not analyze this page: {e}"
            )

    output = (
        "SCANNED PDF VISUAL ANALYSIS\n"
        "===========================\n"
        f"File: {path.name}\n"
        f"Total pages: {total_pages}\n"
        f"Pages analyzed: {pages_to_process}\n"
    )

    if total_pages > pages_to_process:

        output += (
            f"\nOnly the first "
            f"{pages_to_process} pages "
            "were analyzed.\n"
        )

    output += "\n".join(
        analyses
    )

    return output


# ============================================================
# PDF — TEXT + GEMINI FALLBACK
# ============================================================

def _analyze_text_pdf(
    path,
    instruction,
    client,
):
    """
    Complete PDF pipeline.

    Order:

        1. Native Gemini PDF
        2. Scanned visual fallback
        3. Local text extraction
        4. Gemini text analysis
        5. Raw extracted content
    """

    # ========================================================
    # 1. PRIMARY — GEMINI NATIVE PDF
    # ========================================================

    if client is not None:

        direct_result = (
            _analyze_pdf_direct(
                path=path,
                instruction=instruction,
                client=client,
            )
        )

        if (
            direct_result
            and len(
                direct_result.strip()
            ) > 40
        ):

            return direct_result

    # ========================================================
    # 2. CHECK LOCAL TEXT
    # ========================================================

    content = _read_pdf(
        path
    )

    has_real_text = (
        _pdf_has_selectable_text(
            path
        )
    )

    # ========================================================
    # 3. SCANNED / IMAGE PDF
    # ========================================================

    if not has_real_text:

        print(
            "[FILE PROCESSOR] "
            "No useful selectable PDF text found."
        )

        if client is not None:

            print(
                "[FILE PROCESSOR] "
                "Starting Gemini visual page fallback..."
            )

            visual_result = (
                _analyze_scanned_pdf(
                    path=path,
                    instruction=instruction,
                    client=client,
                )
            )

            if (
                visual_result
                and "PDF ANALYSIS FAILED"
                not in visual_result
            ):

                return visual_result

        return (
            "PDF ANALYSIS FAILED\n"
            "===================\n"
            f"File: {path.name}\n\n"

            "No readable text was extracted and "
            "Gemini visual analysis was unavailable "
            "or unsuccessful.\n\n"

            "The PDF may be corrupted, encrypted, "
            "or its pages may not be renderable."
        )

    # ========================================================
    # 4. GEMINI UNAVAILABLE BUT LOCAL TEXT EXISTS
    # ========================================================

    if client is None:

        return (
            "PDF CONTENT FOR JEEV\n"
            "====================\n"
            f"Name: {path.name}\n"
            "Type: PDF document\n\n"
            f"{_limit_text(content)}"
        )

    # ========================================================
    # 5. GEMINI TEXT ANALYSIS
    # ========================================================

    try:

        prompt_instruction = (
            instruction or ""
        ).strip()

        if not prompt_instruction:

            prompt_instruction = (
                "Analyze this PDF document carefully. "

                "If it is a CV/resume, provide a "
                "structured professional analysis covering:\n"

                "1. Candidate name and profile\n"
                "2. Education\n"
                "3. Work experience\n"
                "4. Internships\n"
                "5. Skills\n"
                "6. Projects\n"
                "7. Certifications\n"
                "8. Achievements\n"
                "9. Languages\n"
                "10. Strengths\n"
                "11. Weaknesses or missing information\n"
                "12. Resume quality and structure\n"
                "13. Suggestions for improvement\n"
                "14. Overall assessment\n\n"

                "Do not invent information."
            )

        prompt = (
            f"Analyze the extracted content "
            f"from '{path.name}'.\n\n"

            f"{prompt_instruction}\n\n"

            "DOCUMENT CONTENT\n"
            "================\n"

            f"{_limit_text(content)}"
        )

        response = (
            client.models.generate_content(
                model=_get_file_model(),
                contents=prompt,
            )
        )

        result = getattr(
            response,
            "text",
            None,
        )

        if result and result.strip():

            return (
                "PDF ANALYSIS\n"
                "============\n"
                f"File: {path.name}\n\n"
                f"{result.strip()}"
            )

    except Exception as e:

        print(
            "[FILE PROCESSOR] "
            f"Extracted-text Gemini analysis failed: {e}"
        )

    # ========================================================
    # 6. RAW CONTENT FALLBACK
    # ========================================================

    return (
        "PDF CONTENT FOR JEEV\n"
        "====================\n"
        f"Name: {path.name}\n"
        "Type: PDF document\n\n"
        f"{_limit_text(content)}"
    )


# ============================================================
# DOCX
# ============================================================

def _read_docx(path):

    try:
        from docx import Document

    except ImportError:
        return (
            "DOCX support is not installed.\n"
            "Install it with:\n"
            "python -m pip install python-docx"
        )

    try:

        document = Document(
            str(path)
        )

        content = []

        for paragraph in document.paragraphs:

            text = (
                paragraph.text
                .strip()
            )

            if text:
                content.append(text)

        for table_index, table in enumerate(
            document.tables
        ):

            content.append(
                f"\n--- Table "
                f"{table_index + 1} ---"
            )

            for row in table.rows:

                values = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                content.append(
                    " | ".join(values)
                )

        return "\n".join(
            content
        )

    except Exception as e:

        return (
            f"DOCX processing failed: {e}"
        )


# ============================================================
# XLSX
# ============================================================

def _read_xlsx(path):

    try:
        import openpyxl

    except ImportError:
        return (
            "Excel support is not installed.\n"
            "Install it with:\n"
            "python -m pip install openpyxl"
        )

    workbook = None

    try:

        workbook = (
            openpyxl.load_workbook(
                filename=str(path),
                read_only=True,
                data_only=True,
            )
        )

        output = []

        for sheet in workbook.worksheets:

            output.append(
                f"\n=== Sheet: "
                f"{sheet.title} ==="
            )

            row_count = 0

            for row in sheet.iter_rows(
                values_only=True
            ):

                if (
                    row_count
                    >= MAX_SHEET_ROWS
                ):

                    output.append(
                        "[Remaining rows omitted.]"
                    )

                    break

                values = []

                for value in row:

                    if value is None:
                        values.append("")

                    else:

                        text = str(
                            value
                        )

                        if (
                            len(text)
                            > MAX_CELL_CHARS
                        ):

                            text = (
                                text[
                                    :MAX_CELL_CHARS
                                ]
                                + "..."
                            )

                        values.append(text)

                output.append(
                    " | ".join(values)
                )

                row_count += 1

        return "\n".join(
            output
        )

    except Exception as e:

        return (
            f"Excel processing failed: {e}"
        )

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass


# ============================================================
# CSV
# ============================================================

def _read_csv(path):

    try:

        output = []

        with open(
            path,
            "r",
            encoding="utf-8",
            errors="replace",
            newline="",
        ) as file:

            reader = csv.reader(
                file
            )

            for index, row in enumerate(
                reader
            ):

                if (
                    index
                    >= MAX_SHEET_ROWS
                ):

                    output.append(
                        "[Remaining rows omitted.]"
                    )

                    break

                output.append(
                    " | ".join(row)
                )

        return "\n".join(
            output
        )

    except Exception as e:

        return (
            f"CSV processing failed: {e}"
        )


# ============================================================
# PPTX
# ============================================================

def _read_pptx(path):

    try:
        from pptx import Presentation

    except ImportError:
        return (
            "PowerPoint support is not installed.\n"
            "Install it with:\n"
            "python -m pip install python-pptx"
        )

    try:

        presentation = (
            Presentation(
                str(path)
            )
        )

        output = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):

            output.append(
                f"\n--- Slide "
                f"{slide_number} ---"
            )

            for shape in slide.shapes:

                if not hasattr(
                    shape,
                    "text",
                ):
                    continue

                text = (
                    shape.text
                    .strip()
                )

                if text:
                    output.append(text)

        return "\n".join(
            output
        )

    except Exception as e:

        return (
            f"PowerPoint processing failed: {e}"
        )


# ============================================================
# FILE TYPE
# ============================================================

def _file_type(path):

    extension = (
        path.suffix.lower()
    )

    if extension == ".pdf":
        return "PDF document"

    if extension == ".docx":
        return "Word document"

    if extension in (
        ".xlsx",
        ".xls",
    ):
        return "Excel spreadsheet"

    if extension == ".pptx":
        return "PowerPoint presentation"

    if extension in IMAGE_EXTENSIONS:
        return "Image"

    if extension in TEXT_EXTENSIONS:
        return "Text/code document"

    if extension == ".csv":
        return "CSV spreadsheet"

    if extension == ".json":
        return "JSON document"

    mime, _ = mimetypes.guess_type(
        str(path)
    )

    if mime:
        return mime

    return (
        extension[1:].upper()
        + " file"
        if extension
        else "Unknown file"
    )


# ============================================================
# CONTENT EXTRACTION
# ============================================================

def _extract_content(path):

    extension = (
        path.suffix.lower()
    )

    if extension in TEXT_EXTENSIONS:
        return _read_text_file(path)

    if extension == ".pdf":
        return _read_pdf(path)

    if extension == ".docx":
        return _read_docx(path)

    if extension in (
        ".xlsx",
        ".xls",
    ):
        return _read_xlsx(path)

    if extension == ".csv":
        return _read_csv(path)

    if extension == ".pptx":
        return _read_pptx(path)

    if extension in IMAGE_EXTENSIONS:

        return (
            "This is an image file.\n\n"
            + _image_info(path)
            + "\n\n"
            "Use the image_analysis action "
            "to understand the image."
        )

    return (
        f"No text extractor is configured "
        f"for '{extension or 'this file type'}'."
    )


# ============================================================
# IMAGE ANALYSIS
# ============================================================

def _analyze_image(
    path,
    instruction,
    client,
):

    if client is None:

        return (
            "The image was received, but JEEV's "
            "Gemini vision client is not available "
            "right now."
        )

    try:

        file_size = path.stat().st_size

        if file_size > IMAGE_MAX_BYTES:

            return (
                "The image is too large for "
                "direct analysis "
                f"({file_size / (1024 * 1024):.1f} MB).\n"

                "Maximum supported size here is "
                f"{IMAGE_MAX_BYTES / (1024 * 1024):.0f} MB."
            )

        mime_type, _ = (
            mimetypes.guess_type(
                str(path)
            )
        )

        if (
            not mime_type
            or not mime_type.startswith(
                "image/"
            )
        ):
            mime_type = "image/jpeg"

        with open(
            path,
            "rb",
        ) as image_file:

            image_bytes = (
                image_file.read()
            )

        from google.genai import types

        image_part = (
            types.Part.from_bytes(
                data=image_bytes,
                mime_type=mime_type,
            )
        )

        prompt = (
            instruction or ""
        ).strip()

        if not prompt:

            prompt = (
                "Analyze this image carefully. "

                "Describe what is visible, identify "
                "important objects, people, text, "
                "symbols, layout and other useful details. "

                "If text is visible, transcribe the "
                "important text. "

                "Be accurate and do not invent details."
            )

        response = (
            client.models.generate_content(
                model=_get_file_model(),
                contents=[
                    image_part,
                    prompt,
                ],
            )
        )

        result = getattr(
            response,
            "text",
            None,
        )

        if result:

            return (
                "IMAGE ANALYSIS\n"
                "==============\n"
                f"File: {path.name}\n\n"
                f"{result.strip()}"
            )

        return (
            "Gemini received the image but "
            "returned no text analysis."
        )

    except Exception as e:

        return (
            f"Image analysis failed: {e}"
        )


# ============================================================
# PDF ANALYSIS ENTRY POINT
# ============================================================

def _process_pdf(
    path,
    action,
    instruction,
    client,
):

    pdf_analysis_actions = {
        "pdf_analysis",
        "analyze_pdf",
        "analyse_pdf",
        "cv_analysis",
        "resume_analysis",
        "analyze_cv",
        "analyse_cv",
        "analyze_resume",
        "analyse_resume",
    }

    if action in pdf_analysis_actions:

        return _analyze_text_pdf(
            path=path,
            instruction=instruction,
            client=client,
        )

    if action in (
        "analyze",
        "analyse",
        "summarize",
        "summary",
        "summary_text",
    ):

        return _analyze_text_pdf(
            path=path,
            instruction=instruction,
            client=client,
        )

    if action in (
        "read",
        "extract",
        "content",
    ):

        content = _read_pdf(
            path
        )

        return (
            "FILE CONTENT FOR JEEV\n"
            "=====================\n"
            f"Name: {path.name}\n"
            "Type: PDF document\n"
            f"Path: {path}\n\n"
            f"{_limit_text(content)}"
        )

    content = _read_pdf(
        path
    )

    return (
        "FILE CONTENT FOR JEEV\n"
        "=====================\n"
        f"Name: {path.name}\n"
        "Type: PDF document\n"
        f"Path: {path}\n\n"
        f"{_limit_text(content)}"
    )


# ============================================================
# PERSISTENT DOCUMENT MEMORY
# ============================================================

def _remember_analysis(
    path,
    result,
    file_type,
    action,
):
    """
    Persist useful file/PDF analysis.

    The complete result is passed to the memory manager.
    """

    if save_document_memory is None:
        print(
            "[FILE PROCESSOR] "
            "Document memory unavailable."
        )
        return

    if not result:
        return

    text = str(
        result
    ).strip()

    if len(text) < 40:
        return

    # Don't store metadata-only operations.
    if action in {
        "info",
        "word_count",
        "json",
    }:
        return

    try:

        summary = text

        if len(summary) > 1800:

            summary = (
                summary[:1797]
                .rstrip()
                + "..."
            )

        save_document_memory(
            file_name=path.name,
            file_path=str(path),
            content=text,
            file_type=file_type,
            summary=summary,
        )

        print(
            "[FILE PROCESSOR] "
            f"Document memory saved: "
            f"{path.name}"
        )

    except Exception as e:

        print(
            "[FILE PROCESSOR] "
            f"Memory save warning: {e}"
        )


# ============================================================
# MAIN FILE PROCESSOR
# ============================================================

def file_processor(
    parameters=None,
    player=None,
    speak=None,
    client=None,
):

    parameters = (
        parameters
        or {}
    )

    file_path = _safe_string(
        parameters.get(
            "file_path",
            "",
        )
    )

    action = _safe_string(
        parameters.get(
            "action",
            "info",
        )
    ).lower()

    instruction = _safe_string(
        parameters.get(
            "instruction",
            "",
        )
    )

    if not file_path:
        return "No file was provided."

    path = Path(
        file_path
    )

    if not path.exists():
        return (
            f"File not found: "
            f"{file_path}"
        )

    if not path.is_file():
        return (
            f"That path is not a file: "
            f"{file_path}"
        )

    try:

        file_size = (
            path.stat().st_size
        )

        extension = (
            path.suffix.lower()
        )

        file_type = _file_type(
            path
        )

        # ====================================================
        # LOCATE GEMINI CLIENT
        # ====================================================

        gemini_client = (
            _find_gemini_client(
                player=player,
                client=client,
            )
        )

        if gemini_client is not None:

            print(
                "[FILE PROCESSOR] "
                "Gemini client available."
            )

        else:

            print(
                "[FILE PROCESSOR] "
                "WARNING: Gemini client unavailable."
            )

        # ====================================================
        # PDF
        # ====================================================

        if extension == ".pdf":

            result = _process_pdf(
                path=path,
                action=action,
                instruction=instruction,
                client=gemini_client,
            )

            _remember_analysis(
                path,
                result,
                file_type,
                action,
            )

            return result

        # ====================================================
        # IMAGE ANALYSIS
        # ====================================================

        if (
            action in (
                "image_analysis",
                "image_analyze",
                "vision",
                "visual_analysis",
                "describe_image",
                "analyze_image",
                "analyse_image",
            )
            or (
                extension
                in IMAGE_EXTENSIONS
                and action
                in (
                    "analyze",
                    "analyse",
                    "read",
                    "describe",
                )
            )
        ):

            if extension not in IMAGE_EXTENSIONS:

                return (
                    f"'{path.name}' "
                    "is not a supported image file."
                )

            result = _analyze_image(
                path=path,
                instruction=instruction,
                client=gemini_client,
            )

            _remember_analysis(
                path,
                result,
                file_type,
                action,
            )

            return result

        # ====================================================
        # INFO
        # ====================================================

        if action == "info":

            return (
                f"File: {path.name}\n"
                f"Path: {path}\n"
                f"Type: {file_type}\n"
                f"Extension: "
                f"{path.suffix or 'none'}\n"
                f"Size: {file_size:,} bytes"
            )

        # ====================================================
        # READ / EXTRACT
        # ====================================================

        if action in (
            "read",
            "extract",
            "content",
        ):

            content = _extract_content(
                path
            )

            return (
                "FILE CONTENT FOR JEEV\n"
                "======================\n"
                f"Name: {path.name}\n"
                f"Type: {file_type}\n"
                f"Path: {path}\n\n"
                f"{_limit_text(content)}"
            )

        # ====================================================
        # ANALYZE NON-PDF FILE
        # ====================================================

        if action in (
            "analyze",
            "analyse",
            "summarize",
            "summary",
            "summary_text",
        ):

            content = _extract_content(
                path
            )

            if gemini_client is not None:

                try:

                    prompt = (
                        instruction
                        or
                        (
                            f"Analyze the following "
                            f"{file_type} named "
                            f"'{path.name}'. "

                            "Extract and explain the "
                            "important information accurately. "

                            "Do not invent information.\n\n"

                            "FILE CONTENT\n"
                            "============\n"

                            f"{_limit_text(content)}"
                        )
                    )

                    response = (
                        gemini_client
                        .models
                        .generate_content(
                            model=_get_file_model(),
                            contents=prompt,
                        )
                    )

                    result = getattr(
                        response,
                        "text",
                        None,
                    )

                    if result:

                        final_result = (
                            "FILE ANALYSIS\n"
                            "=============\n"
                            f"File: {path.name}\n\n"
                            f"{result.strip()}"
                        )

                        _remember_analysis(
                            path,
                            final_result,
                            file_type,
                            action,
                        )

                        return final_result

                except Exception as e:

                    print(
                        "[FILE PROCESSOR] "
                        "Non-PDF Gemini analysis failed: "
                        f"{e}"
                    )

            fallback_result = (
                "FILE CONTENT FOR JEEV\n"
                "======================\n"
                f"Name: {path.name}\n"
                f"Type: {file_type}\n"
                f"Path: {path}\n\n"
                f"{_limit_text(content)}"
            )

            _remember_analysis(
                path,
                fallback_result,
                file_type,
                action,
            )

            return fallback_result

        # ====================================================
        # WORD COUNT
        # ====================================================

        if action == "word_count":

            content = _extract_content(
                path
            )

            if not content:
                return (
                    "No readable text was found."
                )

            return (
                f"Word count: "
                f"{len(content.split()):,}"
            )

        # ====================================================
        # JSON
        # ====================================================

        if action == "json":

            if extension != ".json":
                return (
                    "The selected file is not JSON."
                )

            text = _read_text_file(
                path
            )

            try:

                parsed = json.loads(
                    text
                )

                return json.dumps(
                    parsed,
                    indent=2,
                    ensure_ascii=False,
                )[
                    :MAX_TEXT_CHARS
                ]

            except Exception as e:

                return (
                    f"Invalid JSON: {e}"
                )

        # ====================================================
        # UNKNOWN ACTION
        # ====================================================

        return (
            f"File processor received "
            f"action '{action}' for "
            f"'{path.name}', but that "
            "operation is not implemented."
        )

    except Exception as e:

        print(
            "[FILE PROCESSOR] "
            f"Fatal processing error: {e}"
        )

        return (
            f"File processing failed: {e}"
        )